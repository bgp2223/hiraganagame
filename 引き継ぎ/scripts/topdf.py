import openpyxl, datetime, html
from openpyxl.utils import get_column_letter as L
orig=openpyxl.load_workbook('f6.xlsx',data_only=True)['R8表紙']
new=openpyxl.load_workbook('f6.xlsx')['R8表紙']
MAXR,MAXC=117,18
def val(r,c):
    v=new.cell(r,c).value
    if isinstance(v,str) and v.startswith('='): return orig.cell(r,c).value
    return v
def fmt(v,nf):
    if v is None: return ''
    if isinstance(v,(datetime.datetime,datetime.date)):
        n=nf.lower()
        if nf=='mm-dd-yy': return f'{v.year}/{v.month}/{v.day}'
        if 'm' in n and 'd' in n: return f'{v.month}/{v.day}'
        if 'd' in n: return str(v.day)
        return f'{v.month}/{v.day}'
    if 'aaa' in nf and isinstance(v,(int,float)):
        return '月火水木金土日'[(datetime.date(1899,12,30)+datetime.timedelta(days=int(v))).weekday()]
    if isinstance(v,float): return str(int(v)) if v==int(v) else f'{v:.1f}'
    return str(v)
skip=set(); span={}
for m in new.merged_cells.ranges:
    if m.min_row>MAXR or m.min_col>MAXC: continue
    span[(m.min_row,m.min_col)]=(min(m.max_row,MAXR)-m.min_row+1,min(m.max_col,MAXC)-m.min_col+1)
    for r in range(m.min_row,min(m.max_row,MAXR)+1):
        for c in range(m.min_col,min(m.max_col,MAXC)+1):
            if (r,c)!=(m.min_row,m.min_col): skip.add((r,c))
bs={'thin':'1px solid #000','medium':'2px solid #000','double':'3px double #000','hair':'1px dotted #000','dotted':'1px dotted #000','dashed':'1px dashed #000','thick':'3px solid #000'}
def bd(b): return bs.get(b.style,'') if b and b.style else ''
def color(cobj):
    try:
        if cobj is not None and cobj.type=='rgb' and isinstance(cobj.rgb,str) and cobj.rgb not in('00000000',): return '#'+cobj.rgb[-6:]
    except Exception: pass
    return None
out=['<html><head><meta charset=utf-8><style>@page{size:A3 portrait;margin:6mm}body{margin:0;font-family:"IPAGothic","WenQuanYi Zen Hei",sans-serif}table{border-collapse:collapse;table-layout:fixed}td{padding:0 1px;overflow:hidden;white-space:pre-wrap;line-height:1.15;box-sizing:border-box}</style></head><body><table>']
out.append('<colgroup>'+''.join(f'<col style="width:{(new.column_dimensions[L(c)].width or 8.43)*7+5:.0f}px">' for c in range(1,MAXC+1))+'</colgroup>')
for r in range(1,MAXR+1):
    if r==74: out.append('</table><table style="page-break-before:always">'+out[1])
    h=new.row_dimensions[r].height or 15
    out.append(f'<tr style="height:{h*1.33:.0f}px">')
    for c in range(1,MAXC+1):
        if (r,c) in skip: continue
        cell=new.cell(r,c); v=fmt(val(r,c),cell.number_format)
        st=[]
        f=cell.fill
        if f and f.fill_type=='solid':
            col=color(f.fgColor)
            if col and col!='#FFFFFF': st.append(f'background:{col}')
        fc=color(cell.font.color) if cell.font else None
        if fc: st.append(f'color:{fc}')
        st.append(f'font-size:{(cell.font.sz or 11)*1.2:.1f}px')
        if cell.font.b: st.append('font-weight:bold')
        a=cell.alignment
        st.append('text-align:'+{'center':'center','right':'right','centerContinuous':'center'}.get(a.horizontal,'left'))
        st.append('vertical-align:'+{'top':'top','center':'middle'}.get(a.vertical,'bottom'))
        for side,k in (('left','border-left'),('right','border-right'),('top','border-top'),('bottom','border-bottom')):
            b=bd(getattr(cell.border,side))
            if b: st.append(f'{k}:{b}')
        rs,cs=span.get((r,c),(1,1))
        if v and not a.wrap_text and (r,c) not in span and a.horizontal in (None,'left','general') and not isinstance(val(r,c),(int,float,datetime.datetime)):
            k=c+1
            while k<=MAXC and (r,k) not in skip and (r,k) not in span and val(r,k) is None and k-c<14:
                skip.add((r,k)); cs+=1; k+=1
        if not a.wrap_text: st.append('white-space:nowrap')
        out.append(f'<td rowspan={rs} colspan={cs} style="{";".join(st)}">{html.escape(v)}</td>')
    out.append('</tr>')
out.append('</table></body></html>')
open('sheet.html','w').write(''.join(out))
